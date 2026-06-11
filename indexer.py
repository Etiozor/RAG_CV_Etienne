"""Document indexer for CV Chatbot.

Loads documents from docs/, chunks them, generates embeddings,
and persists to ChromaDB for RAG retrieval.

Supported formats: PDF, Word (.docx), Excel (.xlsx)
"""

import os
from pathlib import Path

from docx import Document as DocxDocument
from dotenv import load_dotenv
from langchain.schema import Document
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.vectorstores import Chroma
from langchain_openai import OpenAIEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook


def load_pdf(file_path: Path) -> list[Document]:
    """Load a PDF file using PyPDFLoader."""
    loader = PyPDFLoader(str(file_path))
    return loader.load()


def load_docx(file_path: Path) -> list[Document]:
    """Load a Word document using python-docx."""
    doc = DocxDocument(str(file_path))

    # Extract all paragraphs
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    content = "\n\n".join(paragraphs)

    return [Document(
        page_content=content,
        metadata={"source": str(file_path), "type": "docx"}
    )]


def load_xlsx(file_path: Path) -> list[Document]:
    """Load an Excel file using openpyxl."""
    wb = load_workbook(str(file_path), data_only=True)
    documents = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = []

        for row in sheet.iter_rows(values_only=True):
            # Filter out empty cells and join with separator
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                rows.append(" | ".join(cells))

        if rows:
            content = "\n".join(rows)
            documents.append(Document(
                page_content=content,
                metadata={"source": str(file_path), "sheet": sheet_name, "type": "xlsx"}
            ))

    return documents


def load_documents(docs_dir: Path) -> list[Document]:
    """Load all documents from the specified directory.

    Supports: PDF, Word (.docx), Excel (.xlsx)

    Args:
        docs_dir: Path to the documents directory

    Returns:
        List of LangChain Document objects
    """
    documents = []

    if not docs_dir.exists():
        print(f"Warning: Directory {docs_dir} does not exist")
        return documents

    # Load PDFs
    for pdf_file in docs_dir.glob("*.pdf"):
        try:
            docs = load_pdf(pdf_file)
            documents.extend(docs)
            print(f"  Loaded PDF: {pdf_file.name} ({len(docs)} pages)")
        except Exception as e:
            print(f"  Error loading {pdf_file.name}: {e}")

    # Load Word documents
    for docx_file in docs_dir.glob("*.docx"):
        try:
            docs = load_docx(docx_file)
            documents.extend(docs)
            print(f"  Loaded DOCX: {docx_file.name}")
        except Exception as e:
            print(f"  Error loading {docx_file.name}: {e}")

    # Load Excel files
    for xlsx_file in docs_dir.glob("*.xlsx"):
        try:
            docs = load_xlsx(xlsx_file)
            documents.extend(docs)
            print(f"  Loaded XLSX: {xlsx_file.name} ({len(docs)} sheets)")
        except Exception as e:
            print(f"  Error loading {xlsx_file.name}: {e}")

    return documents


def chunk_documents(documents: list[Document]) -> list[Document]:
    """Split documents into semantic chunks.

    Uses RecursiveCharacterTextSplitter with:
    - chunk_size: 2000 characters (~500 tokens)
    - chunk_overlap: 200 characters (~50 tokens)

    Args:
        documents: List of Document objects to chunk

    Returns:
        List of chunked Document objects
    """
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=2000,
        chunk_overlap=200,
        length_function=len,
        separators=["\n\n", "\n", " ", ""]
    )

    return text_splitter.split_documents(documents)


def create_vectorstore(
    chunks: list[Document],
    persist_directory: str = "./chroma_db",
    collection_name: str = "cv_documents"
) -> Chroma:
    """Create and persist a ChromaDB vector store.

    Args:
        chunks: List of Document chunks to embed and store
        persist_directory: Path to ChromaDB storage
        collection_name: Name of the collection

    Returns:
        Chroma vector store instance
    """
    # Load environment variables for API key
    load_dotenv()

    if not os.getenv("OPENAI_API_KEY"):
        raise ValueError("OPENAI_API_KEY not found in environment. Copy .env.example to .env and add your key.")

    # Initialize embeddings
    embeddings = OpenAIEmbeddings(model="text-embedding-3-small")

    # Create vector store with persistence
    vectorstore = Chroma.from_documents(
        documents=chunks,
        embedding=embeddings,
        persist_directory=persist_directory,
        collection_name=collection_name
    )

    return vectorstore


def main():
    """Main indexing pipeline."""
    DOCS_DIR = Path("docs")
    CHROMA_DIR = "./chroma_db"

    print("=" * 50)
    print("CV Chatbot - Document Indexer")
    print("=" * 50)

    # Step 1: Load documents
    print("\n[1/3] Loading documents from docs/...")
    documents = load_documents(DOCS_DIR)

    if not documents:
        print("No documents found in docs/. Add your CV, mandates, and profile documents.")
        return

    print(f"\nLoaded {len(documents)} document sections")

    # Step 2: Chunk documents
    print("\n[2/3] Chunking documents...")
    chunks = chunk_documents(documents)
    print(f"Created {len(chunks)} chunks")

    # Step 3: Create vector store
    print("\n[3/3] Creating embeddings and storing in ChromaDB...")
    vectorstore = create_vectorstore(chunks, CHROMA_DIR)

    print(f"\nIndexed {len(chunks)} chunks to ChromaDB at {CHROMA_DIR}")
    print("=" * 50)
    print("Indexing complete! Ready for RAG queries.")
    print("=" * 50)


if __name__ == "__main__":
    main()
